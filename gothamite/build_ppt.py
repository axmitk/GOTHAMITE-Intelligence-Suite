from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

prs = Presentation()

# Slide 1
slide = prs.slides.add_slide(prs.slide_layouts[0])
title = slide.shapes.title
subtitle = slide.placeholders[1]
title.text = "SMART INDIA HACKATHON 2026"
subtitle.text = "GOTHAMITE\nCross-Source Dark Web Threat Actor De-Anonymization Platform\n\nProblem Statement ID: SIH26151\nProblem Statement Title: Dark Web Threat Actor De-anonymization\nTheme: Blockchain & Cybersecurity\nPS Category: Software\nTeam ID: [TO BE ADDED]\nTeam Name: AvivCREW\nInstitute code: [TO BE ADDED]"
for paragraph in subtitle.text_frame.paragraphs:
    paragraph.font.size = Pt(16)
    paragraph.alignment = PP_ALIGN.LEFT

# Slide 2
slide2 = prs.slides.add_slide(prs.slide_layouts[1])
slide2.shapes.title.text = "GOTHAMITE: Proposed Solution"
tf = slide2.placeholders[1].text_frame
tf.text = "Core Idea: An analyst-centric intelligence platform that deterministically correlates threat actors across decentralized dark web sources using cryptographic and temporal evidence."
p = tf.add_paragraph()
p.text = "How it addresses the problem: Moves beyond naive string matching by evaluating shared PGP fingerprints, cryptocurrency wallets, and temporal succession to counter handle rebranding."
p.level = 0
p = tf.add_paragraph()
p.text = "Innovation & Uniqueness:"
p.level = 0
p = tf.add_paragraph()
p.text = "Contradiction-Aware Scoring: Penalizes concurrent activity overlap (-0.30) to filter out typosquatting decoys."
p.level = 1
p = tf.add_paragraph()
p.text = "Evidence Provenance: Every graph edge links to immutable SHA-256 source artifacts for full auditability."
p.level = 1
p = tf.add_paragraph()
p.text = "Ethical & Safe Pipeline: Treats ingested dark web content as inert data to prevent adversarial injection."
p.level = 1

# Slide 3
slide3 = prs.slides.add_slide(prs.slide_layouts[1])
slide3.shapes.title.text = "TECHNICAL APPROACH"
tf3 = slide3.placeholders[1].text_frame
tf3.text = "Technologies Used:"
p = tf3.add_paragraph()
p.text = "Backend & Framework: Python 3.11, FastAPI (RESTful endpoints)"
p.level = 1
p = tf3.add_paragraph()
p.text = "Database: SQLite / SQLAlchemy (Immutable artifact history)"
p.level = 1
p = tf3.add_paragraph()
p.text = "Graph Engine: NetworkX (In-memory topology traversals & entity resolution)"
p.level = 1
p = tf3.add_paragraph()
p.text = "Analyst UI: Streamlit + streamlit-agraph (Interactive investigation workbench)"
p.level = 1
p = tf3.add_paragraph()
p.text = "Methodology & Process Flow:"
p.level = 0
p = tf3.add_paragraph()
p.text = "Unidirectional 4-Layer Architecture: Ingest -> Storage -> Correlation Engine -> Presentation."
p.level = 1
p = tf3.add_paragraph()
p.text = "Decoupled Processing: Strict separation between ingestion (pure storage) and scoring (batch evaluation) ensures deterministic attribution."
p.level = 1

# Slide 4
slide4 = prs.slides.add_slide(prs.slide_layouts[1])
slide4.shapes.title.text = "FEASIBILITY AND VIABILITY"
tf4 = slide4.placeholders[1].text_frame
tf4.text = "Feasibility Analysis:"
p = tf4.add_paragraph()
p.text = "Technical Feasibility: Uses efficient standard relational DBs and in-memory graph algorithms, highly scalable for batch ops without requiring expensive GPU compute."
p.level = 1
p = tf4.add_paragraph()
p.text = "Operational Feasibility: Streamlit dashboard provides a zero-friction, intuitive UI for investigators."
p.level = 1
p = tf4.add_paragraph()
p.text = "Potential Challenges and Risks:"
p.level = 0
p = tf4.add_paragraph()
p.text = "Adversarial Input: Dark web data is inherently hostile and designed to subvert analysis."
p.level = 1
p = tf4.add_paragraph()
p.text = "False Positives: Naive clustering links distinct actors who merely interact."
p.level = 1
p = tf4.add_paragraph()
p.text = "Strategies for Overcoming Challenges:"
p.level = 0
p = tf4.add_paragraph()
p.text = "Inert Data Defense: Strict boundary validation; content is never passed to eval() or LLMs."
p.level = 1
p = tf4.add_paragraph()
p.text = "Interaction Isolation: 'Transacted with' signals are structurally prevented from merging personas."
p.level = 1

# Slide 5
slide5 = prs.slides.add_slide(prs.slide_layouts[1])
slide5.shapes.title.text = "IMPACT AND BENEFITS"
tf5 = slide5.placeholders[1].text_frame
tf5.text = "Target Audience: Intelligence Analysts, Law Enforcement Agencies (NTRO), Cyber Threat Researchers."
p = tf5.add_paragraph()
p.text = "Operational Impact:"
p.level = 0
p = tf5.add_paragraph()
p.text = "Standardized Attribution: Reduces manual bias by computing linkages through objective, arithmetic signal weights."
p.level = 1
p = tf5.add_paragraph()
p.text = "Evidentiary Integrity: Replaces fragile black-box AI guessing with cryptographically backed, legally admissible digital evidence."
p.level = 1
p = tf5.add_paragraph()
p.text = "Strategic Benefits:"
p.level = 0
p = tf5.add_paragraph()
p.text = "Rapidly uncovers threat actor rebranding, infrastructure migration, and platform cross-posting."
p.level = 1
p = tf5.add_paragraph()
p.text = "Visualizes complex decentralized relationships in an interactive, easy-to-explore node-link graph."
p.level = 1

# Slide 6
slide6 = prs.slides.add_slide(prs.slide_layouts[1])
slide6.shapes.title.text = "RESEARCH AND REFERENCES"
tf6 = slide6.placeholders[1].text_frame
tf6.text = "Digital Evidence Standards: Aligning correlation scoring with strict evidentiary tracking (verifiable SHA-256 artifact hashing)."
p = tf6.add_paragraph()
p.text = "Cryptographic Verification: Implementation of deterministic fingerprint and base58 cryptocurrency wallet matching protocols."
p.level = 0
p = tf6.add_paragraph()
p.text = "Synthetic Onion Routing: Leveraged the companion darkweb-sandbox to validate extraction algorithms against simulated Tor hidden services (.onion.mock) ethically."
p.level = 0
p = tf6.add_paragraph()
p.text = "Architectural Benchmarks: Comparative analysis with standard open-source intelligence tools to ensure GOTHAMITE provides superior contradiction-aware accuracy."
p.level = 0

prs.save('C:/Users/aadik/sih 2026/gothamite/ppt/GOTHAMITE_SIH_2026_Final.pptx')
