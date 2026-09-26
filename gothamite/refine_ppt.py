from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

prs = Presentation()

def add_box(slide, text, left, top, width, height, fill_color=RGBColor(41, 128, 185), font_color=RGBColor(255, 255, 255), font_size=14, bold=True):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.color.rgb = RGBColor(0, 0, 0)
    tf = shape.text_frame
    tf.text = text
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    tf.paragraphs[0].font.color.rgb = font_color
    tf.paragraphs[0].font.size = Pt(font_size)
    tf.paragraphs[0].font.bold = bold
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    return shape

def add_arrow(slide, left, top, width, height):
    shape = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(149, 165, 166)
    shape.line.color.rgb = RGBColor(149, 165, 166)
    return shape

def add_right_arrow(slide, left, top, width, height):
    shape = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(149, 165, 166)
    shape.line.color.rgb = RGBColor(149, 165, 166)
    return shape

# Slide 1: Title
slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout for full control
title_box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1.5))
tf = title_box.text_frame
p = tf.paragraphs[0]
p.text = "SMART INDIA HACKATHON 2026"
p.font.bold = True
p.font.size = Pt(36)
p.font.color.rgb = RGBColor(41, 128, 185)
p.alignment = PP_ALIGN.CENTER

sub_box = slide.shapes.add_textbox(Inches(1), Inches(2.5), Inches(8), Inches(1))
tf = sub_box.text_frame
p = tf.paragraphs[0]
p.text = "GOTHAMITE"
p.font.bold = True
p.font.size = Pt(44)
p.alignment = PP_ALIGN.CENTER

desc_box = slide.shapes.add_textbox(Inches(1), Inches(3.5), Inches(8), Inches(3))
tf = desc_box.text_frame
tf.text = "Problem Statement ID: SIH26151\nProblem Statement Title: Dark Web Threat Actor De-anonymization\nTheme: Blockchain & Cybersecurity\nPS Category: Software\n\nTeam ID: [TO BE ADDED]\nTeam Name: AvivCREW\nInstitute code: [TO BE ADDED]"
for p in tf.paragraphs:
    p.font.size = Pt(18)
    p.alignment = PP_ALIGN.CENTER

# Slide 2: Proposed Solution
slide2 = prs.slides.add_slide(prs.slide_layouts[5])
slide2.shapes.title.text = "GOTHAMITE: PROPOSED SOLUTION"

# Pipeline
add_box(slide2, "Dark-Web Sources", Inches(0.5), Inches(2), Inches(1.5), Inches(0.8))
add_right_arrow(slide2, Inches(2.1), Inches(2.3), Inches(0.3), Inches(0.2))
add_box(slide2, "Evidence Collection", Inches(2.5), Inches(2), Inches(1.6), Inches(0.8))
add_right_arrow(slide2, Inches(4.2), Inches(2.3), Inches(0.3), Inches(0.2))
add_box(slide2, "Entity Extraction", Inches(4.6), Inches(2), Inches(1.5), Inches(0.8))
add_right_arrow(slide2, Inches(6.2), Inches(2.3), Inches(0.3), Inches(0.2))
add_box(slide2, "Deterministic Correlation", Inches(6.6), Inches(2), Inches(1.5), Inches(0.8))
add_right_arrow(slide2, Inches(8.2), Inches(2.3), Inches(0.3), Inches(0.2))
add_box(slide2, "Investigator Graph", Inches(8.6), Inches(2), Inches(1.3), Inches(0.8), fill_color=RGBColor(39, 174, 96))

# Differentiators
add_box(slide2, "Contradiction-Aware Scoring", Inches(1), Inches(4), Inches(2.5), Inches(0.6), fill_color=RGBColor(192, 57, 43))
add_box(slide2, "Evidence Provenance", Inches(4), Inches(4), Inches(2.5), Inches(0.6), fill_color=RGBColor(211, 84, 0))
add_box(slide2, "Interaction Isolation", Inches(7), Inches(4), Inches(2.5), Inches(0.6), fill_color=RGBColor(142, 68, 173))

t1 = slide2.shapes.add_textbox(Inches(1), Inches(4.7), Inches(2.5), Inches(1.5)).text_frame
t1.text = "Penalizes concurrent activity overlap (-0.30) to mathematically filter out typosquatting decoys."
t1.paragraphs[0].font.size = Pt(14)
t1.word_wrap = True

t2 = slide2.shapes.add_textbox(Inches(4), Inches(4.7), Inches(2.5), Inches(1.5)).text_frame
t2.text = "Every graph edge links directly to immutable, SHA-256 hashed source artifacts."
t2.paragraphs[0].font.size = Pt(14)
t2.word_wrap = True

t3 = slide2.shapes.add_textbox(Inches(7), Inches(4.7), Inches(2.5), Inches(1.5)).text_frame
t3.text = "'Transacted with' edges are structurally isolated and never promoted to identity links."
t3.paragraphs[0].font.size = Pt(14)
t3.word_wrap = True

# Slide 3: Technical Approach
slide3 = prs.slides.add_slide(prs.slide_layouts[5])
slide3.shapes.title.text = "TECHNICAL APPROACH & ARCHITECTURE"

add_box(slide3, "INGEST\n(FastAPI)", Inches(1), Inches(2), Inches(2), Inches(0.8))
add_arrow(slide3, Inches(1.9), Inches(2.9), Inches(0.2), Inches(0.4))
add_box(slide3, "STORAGE\n(SQLite/SQLAlchemy)", Inches(1), Inches(3.4), Inches(2), Inches(0.8))
add_arrow(slide3, Inches(1.9), Inches(4.3), Inches(0.2), Inches(0.4))
add_box(slide3, "CORRELATION\n(Deterministic/NetworkX)", Inches(1), Inches(4.8), Inches(2), Inches(0.8))
add_arrow(slide3, Inches(1.9), Inches(5.7), Inches(0.2), Inches(0.4))
add_box(slide3, "PRESENTATION\n(Streamlit UI)", Inches(1), Inches(6.2), Inches(2), Inches(0.8), fill_color=RGBColor(39, 174, 96))

# Data flow
add_box(slide3, "Source Data", Inches(4), Inches(2.5), Inches(1.5), Inches(0.5), fill_color=RGBColor(127, 140, 141))
add_right_arrow(slide3, Inches(5.6), Inches(2.65), Inches(0.3), Inches(0.2))
add_box(slide3, "Immutable Artifacts", Inches(6), Inches(2.5), Inches(1.5), Inches(0.5), fill_color=RGBColor(127, 140, 141))

add_box(slide3, "Entities & Identifiers", Inches(4), Inches(4), Inches(1.5), Inches(0.5), fill_color=RGBColor(127, 140, 141))
add_right_arrow(slide3, Inches(5.6), Inches(4.15), Inches(0.3), Inches(0.2))
add_box(slide3, "Scored Relationships", Inches(6), Inches(4), Inches(1.5), Inches(0.5), fill_color=RGBColor(127, 140, 141))

add_box(slide3, "Graph Visualization", Inches(5), Inches(5.5), Inches(2), Inches(0.8), fill_color=RGBColor(127, 140, 141))

# Connecting lines or text
t_flow = slide3.shapes.add_textbox(Inches(4), Inches(1.8), Inches(4), Inches(0.5)).text_frame
t_flow.text = "Data Object Flow:"
t_flow.paragraphs[0].font.bold = True
t_flow.paragraphs[0].font.size = Pt(16)

# Slide 4: Feasibility
slide4 = prs.slides.add_slide(prs.slide_layouts[5])
slide4.shapes.title.text = "FEASIBILITY, RISK & MITIGATION"

# Risks
add_box(slide4, "HOSTILE INPUT", Inches(1), Inches(2), Inches(2.5), Inches(0.8), fill_color=RGBColor(192, 57, 43))
add_right_arrow(slide4, Inches(3.6), Inches(2.3), Inches(0.4), Inches(0.2))
add_box(slide4, "Strict payload validation (Inert Data Defense)", Inches(4.2), Inches(2), Inches(4), Inches(0.8), fill_color=RGBColor(46, 204, 113))

add_box(slide4, "FALSE LINKAGES", Inches(1), Inches(3.2), Inches(2.5), Inches(0.8), fill_color=RGBColor(192, 57, 43))
add_right_arrow(slide4, Inches(3.6), Inches(3.5), Inches(0.4), Inches(0.2))
add_box(slide4, "Contradiction-aware scoring rules", Inches(4.2), Inches(3.2), Inches(4), Inches(0.8), fill_color=RGBColor(46, 204, 113))

add_box(slide4, "INTERACTION != IDENTITY", Inches(1), Inches(4.4), Inches(2.5), Inches(0.8), fill_color=RGBColor(192, 57, 43))
add_right_arrow(slide4, Inches(3.6), Inches(4.7), Inches(0.4), Inches(0.2))
add_box(slide4, "Absolute isolation of interaction edges", Inches(4.2), Inches(4.4), Inches(4), Inches(0.8), fill_color=RGBColor(46, 204, 113))

add_box(slide4, "AUDITABILITY FAILURE", Inches(1), Inches(5.6), Inches(2.5), Inches(0.8), fill_color=RGBColor(192, 57, 43))
add_right_arrow(slide4, Inches(3.6), Inches(5.9), Inches(0.4), Inches(0.2))
add_box(slide4, "SHA-256 artifact provenance & traceability", Inches(4.2), Inches(5.6), Inches(4), Inches(0.8), fill_color=RGBColor(46, 204, 113))

# Slide 5: Impact & Benefits
slide5 = prs.slides.add_slide(prs.slide_layouts[5])
slide5.shapes.title.text = "IMPACT AND BENEFITS"

add_box(slide5, "STANDARDIZED ATTRIBUTION", Inches(0.5), Inches(2), Inches(2.5), Inches(0.6), font_size=12)
tb1 = slide5.shapes.add_textbox(Inches(0.5), Inches(2.7), Inches(2.5), Inches(1)).text_frame
tb1.text = "Arithmetic, explainable linkage signals"
tb1.paragraphs[0].font.size = Pt(14)
tb1.word_wrap = True

add_box(slide5, "AUDITABLE EVIDENCE", Inches(3.5), Inches(2), Inches(2.5), Inches(0.6), font_size=12)
tb2 = slide5.shapes.add_textbox(Inches(3.5), Inches(2.7), Inches(2.5), Inches(1)).text_frame
tb2.text = "SHA-256 provenance and deep traceability"
tb2.paragraphs[0].font.size = Pt(14)
tb2.word_wrap = True

add_box(slide5, "REDUCED MANUAL BIAS", Inches(0.5), Inches(3.5), Inches(2.5), Inches(0.6), font_size=12)
tb3 = slide5.shapes.add_textbox(Inches(0.5), Inches(4.2), Inches(2.5), Inches(1)).text_frame
tb3.text = "Consistent deterministic scoring without LLM hallucinations"
tb3.paragraphs[0].font.size = Pt(14)
tb3.word_wrap = True

add_box(slide5, "INVESTIGATOR VISIBILITY", Inches(3.5), Inches(3.5), Inches(2.5), Inches(0.6), font_size=12)
tb4 = slide5.shapes.add_textbox(Inches(3.5), Inches(4.2), Inches(2.5), Inches(1)).text_frame
tb4.text = "Interactive, graph-based relationship exploration"
tb4.paragraphs[0].font.size = Pt(14)
tb4.word_wrap = True

# Placeholder for Screenshot
shape = slide5.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.5), Inches(2), Inches(3), Inches(3))
shape.fill.solid()
shape.fill.fore_color.rgb = RGBColor(236, 240, 241)
shape.line.color.rgb = RGBColor(189, 195, 199)
tf_pic = shape.text_frame
tf_pic.text = "[SCREENSHOT TO BE ADDED]"
tf_pic.paragraphs[0].alignment = PP_ALIGN.CENTER
tf_pic.paragraphs[0].font.color.rgb = RGBColor(127, 140, 141)

footer_box = slide5.shapes.add_textbox(Inches(1), Inches(6), Inches(8), Inches(1))
f_tf = footer_box.text_frame
f_tf.text = "Cryptographically backed, auditable evidence designed to support rigorous digital investigations."
f_tf.paragraphs[0].font.italic = True
f_tf.paragraphs[0].font.size = Pt(16)
f_tf.paragraphs[0].alignment = PP_ALIGN.CENTER
f_tf.word_wrap = True

# Slide 6: Research & References
slide6 = prs.slides.add_slide(prs.slide_layouts[5])
slide6.shapes.title.text = "RESEARCH AND REFERENCES"

add_box(slide6, "Digital Evidence Standards", Inches(1), Inches(2), Inches(3.5), Inches(0.8), font_size=16)
tr1 = slide6.shapes.add_textbox(Inches(1), Inches(2.9), Inches(3.5), Inches(1.5)).text_frame
tr1.text = "Aligning correlation logic with strict tracking using verbatim SHA-256 artifact hashing."
tr1.paragraphs[0].font.size = Pt(14)
tr1.word_wrap = True

add_box(slide6, "Cryptographic Verification", Inches(5.5), Inches(2), Inches(3.5), Inches(0.8), font_size=16)
tr2 = slide6.shapes.add_textbox(Inches(5.5), Inches(2.9), Inches(3.5), Inches(1.5)).text_frame
tr2.text = "Implementation of deterministic fingerprint and Base58 cryptocurrency wallet validation."
tr2.paragraphs[0].font.size = Pt(14)
tr2.word_wrap = True

add_box(slide6, "Darkweb-Sandbox Validations", Inches(1), Inches(4.5), Inches(3.5), Inches(0.8), font_size=16)
tr3 = slide6.shapes.add_textbox(Inches(1), Inches(5.4), Inches(3.5), Inches(1.5)).text_frame
tr3.text = "Testing against a custom containerized mock Tor network (.onion.mock) to ethically validate entity extraction."
tr3.paragraphs[0].font.size = Pt(14)
tr3.word_wrap = True

add_box(slide6, "Open-Source Intelligence", Inches(5.5), Inches(4.5), Inches(3.5), Inches(0.8), font_size=16)
tr4 = slide6.shapes.add_textbox(Inches(5.5), Inches(5.4), Inches(3.5), Inches(1.5)).text_frame
tr4.text = "Comparative architectural review against OSINT tools to ensure contradiction-aware defense."
tr4.paragraphs[0].font.size = Pt(14)
tr4.word_wrap = True

prs.save('C:/Users/aadik/sih 2026/gothamite/ppt/GOTHAMITE_SIH_2026_Final_v2.pptx')
